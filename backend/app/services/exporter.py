import os
import re
import requests
import tempfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def get_keywords_from_text(text: str) -> str:
    if not text:
        return "technology"
    cleaned = re.sub(r'[^a-zA-Z0-9\s]', '', text.lower())
    words = cleaned.split()
    stop_words = {"and", "or", "the", "a", "of", "with", "to", "in", "for", "on", "at", "by", "an", "is", "are", "was", "were", "be", "been", "being", "have", "has", "had", "do", "does", "did", "but", "if", "then", "else", "when", "where", "why", "how", "what", "who", "which"}
    filtered_words = [w for w in words if w not in stop_words and len(w) > 2]
    if not filtered_words:
        return "technology"
    return ",".join(filtered_words[:3])

def download_image_for_slide(slide_title: str, suggested_visuals: str) -> str:
    """
    Downloads a relevant image from LoremFlickr based on slide title/visual keywords.
    Returns path to temporary file if successful, otherwise None.
    """
    keywords = get_keywords_from_text(slide_title)
    if suggested_visuals:
        vis_keywords = get_keywords_from_text(suggested_visuals)
        keywords = f"{keywords},{vis_keywords}"
    
    url = f"https://loremflickr.com/400/300/{keywords}"
    try:
        r = requests.get(url, timeout=5)
        if r.status_code == 200 and len(r.content) > 1000:
            with tempfile.NamedTemporaryFile(suffix=".jpg", delete=False) as f:
                f.write(r.content)
                return f.name
    except Exception as e:
        print(f"Error downloading image: {e}")
    
    # Fallback to general technology/education image
    try:
        r = requests.get("https://loremflickr.com/400/300/technology,education", timeout=5)
        if r.status_code == 200 and len(r.content) > 1000:
            with tempfile.NamedTemporaryFile(suffix=".jpg", delete=False) as f:
                f.write(r.content)
                return f.name
    except Exception as e:
        print(f"Error downloading fallback image: {e}")
        
    return None


def generate_pptx_deck(slides_data: list, notes_data: list, output_path: str) -> str:
    """
    Generates an award-winning, startup-grade premium PPTX slide deck with speaker notes bound.
    """
    prs = Presentation()
    temp_images = []
    
    # Choose premium theme colors (Accents, cards, and borders)
    DARK_NAVY = RGBColor(10, 15, 30)       # #0a0f1e
    CYAN = RGBColor(14, 165, 233)          # #0ea5e9
    WHITE = RGBColor(255, 255, 255)
    LIGHT_GRAY = RGBColor(226, 232, 240)    # #e2e8f0
    MUTED_GRAY = RGBColor(148, 163, 184)    # #94a3b8
    CARD_BG = RGBColor(20, 28, 55)          # Deep glassmorphic indigo
    CODE_BG = RGBColor(15, 23, 42)          # Sleek modern terminal background
    CODE_GREEN = RGBColor(34, 197, 94)      # Mint green monospace token

    # Match slides with notes by index
    notes_by_index = {n["slide_index"]: n for n in notes_data}

    for i, slide_info in enumerate(slides_data):
        slide_index = slide_info.get("slide_index", i + 1)
        slide_title = slide_info.get("title", f"Slide {slide_index}")
        bullets = slide_info.get("content", [])
        visuals = slide_info.get("suggested_visuals", "")

        # Add slide layout: Blank Layout (index 6) for complete custom spatial control
        slide_layout = prs.slide_layouts[6] 
        slide = prs.slides.add_slide(slide_layout)

        # Draw main canvas dark-navy solid background
        bg_shape = slide.shapes.add_shape(
            1,  # RECTANGLE shape type
            0, 0, prs.slide_width, prs.slide_height
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = DARK_NAVY
        bg_shape.line.fill.background()

        # DESIGN SELECTION: Title slide vs regular layout
        if i == 0:
            # Layout A: Centered Vertical Showcase (Slide 1 Title Deck)
            # Main Title
            title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(8.0), Inches(2.0))
            tf_title = title_box.text_frame
            tf_title.word_wrap = True
            p_title = tf_title.paragraphs[0]
            p_title.text = slide_title
            p_title.alignment = PP_ALIGN.CENTER
            p_title.font.name = "Georgia"
            p_title.font.size = Pt(40)
            p_title.font.bold = True
            p_title.font.color.rgb = CYAN

            # Decorative Accent Rule (Neon line separating title and subtitle)
            rule = slide.shapes.add_shape(
                1, Inches(2.5), Inches(4.0), Inches(5.0), Inches(0.04)
            )
            rule.fill.solid()
            rule.fill.fore_color.rgb = CYAN
            rule.line.fill.background()

            # Subtitle
            sub_box = slide.shapes.add_textbox(Inches(1.0), Inches(4.3), Inches(8.0), Inches(1.5))
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = "INTELLIGENT EDUCATION WORKFLOWS • KALYX"
            p_sub.alignment = PP_ALIGN.CENTER
            p_sub.font.name = "Arial"
            p_sub.font.size = Pt(13)
            p_sub.font.bold = True
            p_sub.font.color.rgb = MUTED_GRAY

        else:
            # Layout B: Left-aligned Content with Visual Highlight Cards
            # Slide Header Title
            title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8.4), Inches(1.0))
            tf_title = title_box.text_frame
            tf_title.word_wrap = True
            p_title = tf_title.paragraphs[0]
            p_title.text = slide_title
            p_title.font.name = "Georgia"
            p_title.font.size = Pt(28)
            p_title.font.bold = True
            p_title.font.color.rgb = CYAN

            # Check if any bullet point represents a block of code (contains standard tokens)
            has_code = any(
                ("def " in b or "import " in b or "print(" in b or "class " in b or "```" in b)
                for b in bullets
            )

            # Left Column (Bulleted concepts or code)
            content_left = Inches(0.8)
            content_width = Inches(5.5)
            if not visuals:
                content_width = Inches(8.4)  # Take full width if no visual sidebar is present

            if has_code:
                # Draw terminal dark card background for coding block
                code_card = slide.shapes.add_shape(
                    1, content_left, Inches(1.6), content_width, Inches(4.8)
                )
                code_card.fill.solid()
                code_card.fill.fore_color.rgb = CODE_BG
                code_card.line.fill.background()

                code_box = slide.shapes.add_textbox(
                    content_left + Inches(0.2), Inches(1.8), content_width - Inches(0.4), Inches(4.4)
                )
                tf_code = code_box.text_frame
                tf_code.word_wrap = True

                for b_idx, bullet_text in enumerate(bullets):
                    p = tf_code.add_paragraph() if b_idx > 0 else tf_code.paragraphs[0]
                    p.text = bullet_text.replace("```", "")
                    p.font.name = "Consolas"
                    p.font.size = Pt(12)
                    p.font.color.rgb = CODE_GREEN
                    p.space_after = Pt(8)
            else:
                # Standard Bullet Layout
                content_box = slide.shapes.add_textbox(content_left, Inches(1.6), content_width, Inches(4.8))
                tf_content = content_box.text_frame
                tf_content.word_wrap = True

                for b_idx, bullet_text in enumerate(bullets):
                    p = tf_content.add_paragraph() if b_idx > 0 else tf_content.paragraphs[0]
                    p.text = f"•  {bullet_text}"
                    p.font.name = "Arial"
                    p.font.size = Pt(15)
                    p.font.color.rgb = LIGHT_GRAY
                    p.space_after = Pt(14)

            # Right Column: Draw Premium Visual Overlay Highlight Card
            if visuals:
                card_left = Inches(6.8)
                card_width = Inches(2.7)

                # Draw solid background shape for highlight card
                hl_card = slide.shapes.add_shape(
                    1, card_left, Inches(1.6), card_width, Inches(4.8)
                )
                hl_card.fill.solid()
                hl_card.fill.fore_color.rgb = CARD_BG
                hl_card.line.fill.background()

                # Add a thin neon structural rule at the top edge of the card
                top_stripe = slide.shapes.add_shape(
                    1, card_left, Inches(1.6), card_width, Inches(0.08)
                )
                top_stripe.fill.solid()
                top_stripe.fill.fore_color.rgb = CYAN
                top_stripe.line.fill.background()

                # Try to download and add real photo
                temp_img = download_image_for_slide(slide_title, visuals)
                if temp_img:
                    temp_images.append(temp_img)
                    try:
                        # Place image at the top of the card (4:3 aspect ratio)
                        slide.shapes.add_picture(temp_img, card_left, Inches(1.75), width=card_width)
                        
                        # Add caption below the picture
                        vis_box = slide.shapes.add_textbox(
                            card_left + Inches(0.1), Inches(3.9), card_width - Inches(0.2), Inches(2.4)
                        )
                        tf_vis = vis_box.text_frame
                        tf_vis.word_wrap = True

                        p_vis_hdr = tf_vis.paragraphs[0]
                        p_vis_hdr.text = "VISUAL COMPANION"
                        p_vis_hdr.font.name = "Arial"
                        p_vis_hdr.font.size = Pt(10)
                        p_vis_hdr.font.bold = True
                        p_vis_hdr.font.color.rgb = CYAN
                        p_vis_hdr.space_after = Pt(4)

                        p_vis_body = tf_vis.add_paragraph()
                        p_vis_body.text = visuals
                        p_vis_body.font.name = "Arial"
                        p_vis_body.font.size = Pt(9.5)
                        p_vis_body.font.italic = True
                        p_vis_body.font.color.rgb = MUTED_GRAY
                    except Exception as img_err:
                        print(f"Failed to place image on slide: {img_err}")
                        temp_img = None

                if not temp_img:
                    # Visual Suggestion Textbox overlay inside the card (Fallback)
                    vis_box = slide.shapes.add_textbox(
                        card_left + Inches(0.2), Inches(1.8), card_width - Inches(0.4), Inches(4.4)
                    )
                    tf_vis = vis_box.text_frame
                    tf_vis.word_wrap = True

                    p_vis_hdr = tf_vis.paragraphs[0]
                    p_vis_hdr.text = "VISUAL COMPANION"
                    p_vis_hdr.font.name = "Arial"
                    p_vis_hdr.font.size = Pt(11)
                    p_vis_hdr.font.bold = True
                    p_vis_hdr.font.color.rgb = CYAN
                    p_vis_hdr.space_after = Pt(10)

                    p_vis_body = tf_vis.add_paragraph()
                    p_vis_body.text = visuals
                    p_vis_body.font.name = "Arial"
                    p_vis_body.font.size = Pt(11)
                    p_vis_body.font.italic = True
                    p_vis_body.font.color.rgb = MUTED_GRAY
                    p_vis_body.space_after = Pt(8)

        # Bind Speaker/Instructor Notes to Slide Context
        note_match = notes_by_index.get(slide_index)
        if note_match:
            notes_slide = slide.notes_slide
            tf_notes = notes_slide.notes_text_frame
            
            talking_points = note_match.get("talking_points", [])
            tips = note_match.get("teaching_tips", "")
            examples = note_match.get("examples", [])

            notes_content = []
            notes_content.append(f"=== SLIDE {slide_index} INSTRUCTOR NOTES ===")
            
            if talking_points:
                notes_content.append("\nTALKING POINTS:")
                for tp in talking_points:
                    notes_content.append(f"- {tp}")
            
            if tips:
                notes_content.append(f"\nTEACHING TIP:\n{tips}")
                
            if examples:
                notes_content.append("\nCLARIFYING EXAMPLES:")
                for ex in examples:
                    notes_content.append(f"- {ex}")

            tf_notes.text = "\n".join(notes_content)

    # Save presentation package
    try:
        dir_name = os.path.dirname(output_path)
        if dir_name:
            os.makedirs(dir_name, exist_ok=True)
        prs.save(output_path)
    finally:
        # Clean up temp images
        for temp_img in temp_images:
            try:
                if os.path.exists(temp_img):
                    os.unlink(temp_img)
            except Exception as cleanup_err:
                print(f"Error cleaning up temp image {temp_img}: {cleanup_err}")
    return output_path


# PDF Generator using fpdf2 for premium classroom packages
from fpdf import FPDF

class KalyxPDF(FPDF):
    def header(self):
        if self.page_no() > 1:
            self.set_font("helvetica", "B", 8)
            self.set_text_color(148, 163, 184) # Muted gray
            self.cell(0, 10, "KALYX CURRICULUM PACKAGE", border=0, align="L")
            self.cell(0, 10, "CONFIDENTIAL & PROPRIETARY", border=0, align="R")
            self.ln(8)
            self.set_draw_color(226, 232, 240) # Slate line
            self.line(10, 18, 200, 18)
            self.ln(5)

    def footer(self):
        if self.page_no() > 1:
            self.set_y(-15)
            self.set_font("helvetica", "I", 8)
            self.set_text_color(148, 163, 184)
            self.cell(0, 10, f"Page {self.page_no()}/{{nb}}", align="C")

def clean_txt(text: str) -> str:
    if not text:
        return ""
    replacements = {
        "\u2018": "'", "\u2019": "'",  # Smart single quotes
        "\u201c": '"', "\u201d": '"',  # Smart double quotes
        "\u2013": "-", "\u2014": "-",  # Dashes
        "\u2022": "*",                 # Bullets
        "\u2026": "...",               # Ellipsis
        "\u00e9": "e", "\u00e1": "a",  # Accents
        "\u2192": "->",                # Arrows
    }
    for k, v in replacements.items():
        text = text.replace(k, v)
    return text.encode("latin-1", errors="ignore").decode("latin-1")

def generate_pdf_package(course_title: str, course_desc: str, slides: list, notes: list, assessments: list, output_path: str) -> str:
    pdf = KalyxPDF()
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(auto=True, margin=15)
    temp_images = []
    
    # ------------------ COVER PAGE ------------------
    pdf.add_page()
    
    # Large colored band at the top
    pdf.set_fill_color(3, 7, 18) # Dark navy background
    pdf.rect(0, 0, 210, 80, "F")
    
    # Title inside the band
    pdf.set_y(25)
    pdf.set_font("helvetica", "B", 24)
    pdf.set_text_color(14, 165, 233) # Teal/Cyan
    pdf.cell(0, 12, clean_txt(course_title.upper()), align="C", ln=True)
    
    pdf.set_font("helvetica", "B", 10)
    pdf.set_text_color(148, 163, 184)
    pdf.cell(0, 10, "GENERATED BY KALYX CURRICULUM INTELLIGENCE PIPELINE", align="C", ln=True)
    
    # Body info
    pdf.set_y(95)
    pdf.set_text_color(30, 41, 59) # Dark slate
    
    pdf.set_font("helvetica", "B", 14)
    pdf.cell(0, 10, "COURSE OVERVIEW & SYLLABUS DESCRIPTION", ln=True)
    pdf.set_draw_color(14, 165, 233)
    pdf.line(10, 105, 75, 105)
    pdf.ln(4)
    
    pdf.set_font("helvetica", "", 11)
    pdf.multi_cell(0, 6, clean_txt(course_desc or "An advanced curriculum plan detailing learning modules, presentation slides, instructor speaker notes, and multi-tier Bloom's taxonomy assessments."))
    pdf.ln(15)
    
    # Metadata card at the bottom
    pdf.set_fill_color(248, 250, 252) # Light gray background
    pdf.rect(10, 140, 190, 60, "F")
    
    pdf.set_y(145)
    pdf.set_x(15)
    pdf.set_font("helvetica", "B", 11)
    pdf.cell(0, 6, "METADATA PROFILE", ln=True)
    pdf.ln(2)
    
    pdf.set_x(15)
    pdf.set_font("helvetica", "", 10)
    pdf.cell(0, 6, f"Package Type: Full Classroom Implementation Package", ln=True)
    pdf.set_x(15)
    pdf.cell(0, 6, f"Generated Slide Count: {len(slides)} fully-illustrated slides", ln=True)
    pdf.set_x(15)
    pdf.cell(0, 6, f"Assessment Bank size: {len(assessments)} cognitive evaluation items", ln=True)
    pdf.set_x(15)
    pdf.cell(0, 6, f"Pedagogical Standard: Bloom's Revised Taxonomy (Tiers 1-6)", ln=True)
    
    # ------------------ SLIDES & NOTES ------------------
    pdf.add_page()
    
    pdf.set_font("helvetica", "B", 16)
    pdf.set_text_color(14, 165, 233)
    pdf.cell(0, 10, "SECTION 1: INSTRUCTIONAL SLIDE DECK & NOTES", ln=True)
    pdf.set_draw_color(14, 165, 233)
    pdf.line(10, 18, 120, 18)
    pdf.ln(8)
    
    notes_map = {n.get("slide_index"): n for n in notes}
    
    for s in slides:
        slide_idx = s.get("slide_index", 1)
        title = s.get("title", f"Slide {slide_idx}")
        content = s.get("content", [])
        visuals = s.get("suggested_visuals", "")
        
        # Check page remaining height to prevent orphan header card
        if pdf.get_y() > 200:
            pdf.add_page()
            
        # Draw Slide Title Header Card
        pdf.set_fill_color(241, 245, 249) # Subtle gray card
        pdf.set_draw_color(226, 232, 240)
        pdf.rect(10, pdf.get_y(), 190, 12, "FD")
        
        pdf.set_font("helvetica", "B", 11)
        pdf.set_text_color(15, 23, 42)
        pdf.cell(0, 12, f"  SLIDE {slide_idx}: {clean_txt(title).upper()}", ln=True)
        pdf.ln(4)
        
        # Bullets
        pdf.set_font("helvetica", "", 10)
        pdf.set_text_color(51, 65, 85)
        for bullet in content:
            pdf.set_x(15)
            pdf.multi_cell(0, 5, f"-  {clean_txt(bullet)}")
            pdf.ln(1)
            
        pdf.ln(2)
        
        # Visual Companion Box
        if visuals:
            temp_img = download_image_for_slide(title, visuals)
            
            if temp_img:
                temp_images.append(temp_img)
                # Estimate height needed for text on the right
                text_lines = len(visuals) // 65 + 1
                estimated_text_height = 9 + (text_lines * 4.5)
                box_height = max(50, estimated_text_height)
                
                x_pos = 15
                y_pos = pdf.get_y()
                
                if y_pos + box_height > 270:
                    pdf.add_page()
                    y_pos = pdf.get_y()
                
                pdf.set_fill_color(240, 249, 255) # Light cyan background
                pdf.set_draw_color(186, 230, 253) # Light cyan border
                pdf.rect(x_pos, y_pos, 180, box_height, "FD")
                
                # Draw the image
                try:
                    img_y = y_pos + (box_height - 45) / 2
                    pdf.image(temp_img, x=18, y=img_y, w=60, h=45)
                    
                    # Draw text next to image
                    pdf.set_y(y_pos + 4)
                    pdf.set_x(82)
                    pdf.set_font("helvetica", "B", 8)
                    pdf.set_text_color(14, 165, 233)
                    pdf.cell(0, 4, "VISUAL COMPANION COMPONENT SUGGESTION:", ln=True)
                    
                    pdf.set_y(y_pos + 9)
                    pdf.set_x(82)
                    pdf.set_font("helvetica", "I", 9)
                    pdf.set_text_color(71, 85, 105)
                    pdf.multi_cell(110, 4.5, clean_txt(visuals))
                except Exception as img_err:
                    print(f"Failed to place image on PDF: {img_err}")
                    temp_img = None
                
                pdf.set_y(y_pos + box_height + 4)

            if not temp_img:
                pdf.set_x(15)
                pdf.set_fill_color(240, 249, 255) # Light cyan background
                pdf.set_draw_color(186, 230, 253) # Light cyan border
                
                # Estimate height needed for visuals text to draw rect
                text_lines = len(visuals) // 80 + 1
                box_height = 8 + (text_lines * 5)
                
                x_pos = 15
                y_pos = pdf.get_y()
                
                if y_pos + box_height > 270:
                    pdf.add_page()
                    y_pos = pdf.get_y()
                
                pdf.rect(x_pos, y_pos, 180, box_height, "FD")
                pdf.set_y(y_pos + 2)
                pdf.set_x(18)
                pdf.set_font("helvetica", "B", 8)
                pdf.set_text_color(14, 165, 233)
                pdf.cell(0, 4, "VISUAL COMPANION COMPONENT SUGGESTION:", ln=True)
                pdf.set_x(18)
                pdf.set_font("helvetica", "I", 9)
                pdf.set_text_color(71, 85, 105)
                pdf.multi_cell(174, 4.5, clean_txt(visuals))
                pdf.set_y(y_pos + box_height + 4)
            
        # Instructor notes match
        note_match = notes_map.get(slide_idx)
        if note_match:
            pdf.ln(2)
            pdf.set_x(15)
            pdf.set_font("helvetica", "B", 9)
            pdf.set_text_color(14, 165, 233)
            pdf.cell(0, 5, "LECTURE SPEAKER NOTES & PEDAGOGY:", ln=True)
            pdf.ln(1)
            
            talking_points = note_match.get("talking_points", [])
            tips = note_match.get("teaching_tips", "")
            examples = note_match.get("examples", [])
            
            if talking_points:
                pdf.set_x(18)
                pdf.set_font("helvetica", "B", 8.5)
                pdf.set_text_color(71, 85, 105)
                pdf.cell(0, 4, "Talking Points:", ln=True)
                pdf.set_font("helvetica", "", 9)
                pdf.set_text_color(51, 65, 85)
                for tp in talking_points:
                    pdf.set_x(20)
                    pdf.multi_cell(0, 4, f"-  {clean_txt(tp)}")
                    pdf.ln(0.5)
            
            if tips:
                pdf.ln(1)
                pdf.set_x(18)
                pdf.set_font("helvetica", "B", 8.5)
                pdf.set_text_color(71, 85, 105)
                pdf.cell(0, 4, "Pedagogical Tip:", ln=True)
                pdf.set_x(20)
                pdf.set_font("helvetica", "I", 9)
                pdf.set_text_color(51, 65, 85)
                pdf.multi_cell(0, 4, clean_txt(tips))
                
            if examples:
                pdf.ln(1)
                pdf.set_x(18)
                pdf.set_font("helvetica", "B", 8.5)
                pdf.set_text_color(71, 85, 105)
                pdf.cell(0, 4, "Classroom Examples:", ln=True)
                pdf.set_font("helvetica", "", 9)
                pdf.set_text_color(51, 65, 85)
                for ex in examples:
                    pdf.set_x(20)
                    pdf.multi_cell(0, 4, f"> \"{clean_txt(ex)}\"")
                    pdf.ln(0.5)
                    
        pdf.ln(10)
        
        # Add a subtle horizontal divider between slides
        if pdf.get_y() < 250:
            pdf.set_draw_color(241, 245, 249)
            pdf.line(10, pdf.get_y(), 200, pdf.get_y())
            pdf.ln(4)
            
    # ------------------ ASSESSMENT BANK ------------------
    pdf.add_page()
    
    pdf.set_font("helvetica", "B", 16)
    pdf.set_text_color(14, 165, 233)
    pdf.cell(0, 10, "SECTION 2: PEDAGOGICAL ASSESSMENT BANK", ln=True)
    pdf.set_draw_color(14, 165, 233)
    pdf.line(10, 18, 120, 18)
    pdf.ln(8)
    
    for a_idx, a in enumerate(assessments):
        q_text = a.get("question_text", "")
        q_type = a.get("question_type", "MCQ")
        options = a.get("options", [])
        correct = a.get("correct_answer", "")
        bloom = a.get("bloom_level", "Remembering")
        
        # Check page remaining height to prevent orphan question card
        if pdf.get_y() > 220:
            pdf.add_page()
            
        # Draw Question Header card
        pdf.set_fill_color(248, 250, 252)
        pdf.rect(10, pdf.get_y(), 190, 8, "F")
        pdf.set_font("helvetica", "B", 10)
        pdf.set_text_color(15, 23, 42)
        pdf.cell(0, 8, f" QUESTION {a_idx + 1} ({q_type.upper()}) [BLOOM TIER: {bloom.upper()}]", ln=True)
        pdf.ln(3)
        
        pdf.set_font("helvetica", "", 10)
        pdf.set_text_color(51, 65, 85)
        pdf.set_x(12)
        pdf.multi_cell(0, 5, clean_txt(q_text))
        pdf.ln(2)
        
        if options and q_type.upper() == "MCQ":
            pdf.set_font("helvetica", "", 9.5)
            for opt_idx, opt in enumerate(options):
                pdf.set_x(16)
                pdf.multi_cell(0, 4, f"({chr(65 + opt_idx)})  {clean_txt(opt)}")
                pdf.ln(1)
            pdf.ln(1)
            
        if correct:
            pdf.set_x(12)
            pdf.set_font("helvetica", "B", 9)
            pdf.set_text_color(22, 163, 74) # Green color for correct answer
            pdf.cell(0, 5, f"Correct Answer:  {clean_txt(correct)}", ln=True)
            
        pdf.ln(6)
        
        # Add page break if question is near bottom
        if pdf.get_y() > 240 and a_idx < len(assessments) - 1:
            pdf.add_page()
            
    # Save PDF
    try:
        dir_name = os.path.dirname(output_path)
        if dir_name:
            os.makedirs(dir_name, exist_ok=True)
        pdf.output(output_path)
    finally:
        # Clean up temp images
        for temp_img in temp_images:
            try:
                if os.path.exists(temp_img):
                    os.unlink(temp_img)
            except Exception as cleanup_err:
                print(f"Error cleaning up temp image {temp_img}: {cleanup_err}")
    return output_path
